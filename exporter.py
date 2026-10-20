"""Безопасная запись CSV и Excel: сначала временный файл, потом подмена; старая версия уходит в backup/."""
import csv
import json
import os
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("title", "Название"),
    ("price", "Цена"),
    ("currency", "Валюта"),
    ("in_stock", "В наличии"),
    ("stock_qty", "Остаток, шт."),
    ("category", "Категория"),
    ("rating", "Рейтинг (1-5)"),
    ("sku", "Артикул (UPC)"),
    ("specs", "Характеристики"),
    ("url", "Ссылка"),
]
WIDTHS = [40, 9, 8, 11, 12, 20, 12, 20, 70, 70]


def _rows(products):
    rows = []
    for p in products:
        try:
            specs = "; ".join(f"{k}: {v}" for k, v in json.loads(p["specs"] or "{}").items())
        except ValueError:
            specs = ""
        rows.append([
            p["title"], p["price"], p["currency"], "да" if p["in_stock"] else "нет",
            p["stock_qty"], p["category"], p["rating"], p["sku"], specs, p["url"],
        ])
    return rows


def _write_csv(path, rows):
    # utf-8-sig и «;» — чтобы русский Excel открыл файл сразу правильно, без мастера импорта
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([title for _, title in COLUMNS])
        writer.writerows(rows)


def _write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append([title for _, title in COLUMNS])
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for line in ws.iter_rows(min_row=2):
        line[1].number_format = "0.00"
        for cell in line:
            # текст, начинающийся с «=», не должен превращаться в формулу
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    wb.save(path)


def _safe_write(target, write_func, backup_dir):
    target = Path(target)
    tmp = target.with_name("~tmp_" + target.name)
    try:
        write_func(tmp)
    except Exception:
        tmp.unlink(missing_ok=True)
        raise
    if target.exists():
        backup_dir.mkdir(exist_ok=True)
        shutil.copy2(target, backup_dir / f"{target.stem}_{datetime.now():%Y%m%d_%H%M%S}{target.suffix}")
    try:
        os.replace(tmp, target)
        return target
    except PermissionError:
        # файл открыт в Excel и заблокирован — сохраняем рядом под другим именем
        alt = target.with_name(f"{target.stem}_{datetime.now():%H%M%S}{target.suffix}")
        os.replace(tmp, alt)
        return alt


def export_all(products, run_dir):
    run_dir = Path(run_dir)
    backup_dir = run_dir / "backup"
    rows = _rows(products)
    return [
        _safe_write(run_dir / "products.csv", lambda p: _write_csv(p, rows), backup_dir),
        _safe_write(run_dir / "products.xlsx", lambda p: _write_xlsx(p, rows), backup_dir),
    ]